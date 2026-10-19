"""Exports a project's final transcript as plain text, SRT subtitles, or docx."""
import io
from typing import Any, Literal

from app.services.transcript import merge_resolved

ExportFormat = Literal["txt", "srt", "docx"]


def _format_srt_timestamp(seconds: float) -> str:
    millis = round(seconds * 1000)
    hours, millis = divmod(millis, 3_600_000)
    minutes, millis = divmod(millis, 60_000)
    secs, millis = divmod(millis, 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"


def build_txt(final_text: str) -> bytes:
    return final_text.encode("utf-8")


def build_srt(
    reference_segments: list[dict[str, Any]],
    resolved_segments: list[dict[str, Any]] | None = None,
    max_words_per_cue: int = 12,
) -> bytes:
    # Splice in human-resolved corrections so the SRT matches the final_text
    # export instead of the raw (possibly doubtful) reference transcript.
    items = merge_resolved(reference_segments, resolved_segments or [])
    lines: list[str] = []
    cue_number = 1
    i = 0
    while i < len(items):
        chunk = items[i : i + max_words_per_cue]
        start = chunk[0]["start"]
        end = chunk[-1]["end"]
        text = " ".join(item["text"] for item in chunk if item["text"])
        lines.append(str(cue_number))
        lines.append(f"{_format_srt_timestamp(start)} --> {_format_srt_timestamp(end)}")
        lines.append(text)
        lines.append("")
        cue_number += 1
        i += max_words_per_cue
    return "\n".join(lines).encode("utf-8")


def build_docx(final_text: str, title: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    for paragraph in final_text.split("\n") or [final_text]:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


CONTENT_TYPES = {
    "txt": "text/plain",
    "srt": "application/x-subrip",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
